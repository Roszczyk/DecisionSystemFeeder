
import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INPUT_FILE = Path(__file__).parent / "output.log"
OUTPUT_FILE = Path(__file__).parent / "fusion_results_comparison.xlsx"

METRICS = [
    "including correct",
    "only correct",
    "correct parts",
]

METHODS = [
    "simple voting",
    "cumulative voting",
    "approval voting",
    "borda voting",
    "nanson voting",
    "DST belief interval with belief focus",
    "DST belief interval with plausibility focus",
    "DST belief",
    "DST plausibility",
    "DST mass",
]

SCENARIOS = [
    "Low uncertainty",
    "High uncertainty",
]


def parse_output(text):
    """Extract TOTAL RESULTS metrics from the simulation log."""
    results = {}
    scenario = None
    method = None
    in_total_results = False

    for line in text.splitlines():
        line = line.strip()

        if line.startswith("Scenario 1: Low uncertainty"):
            scenario = "Low uncertainty"
            results[scenario] = {}
            method = None
            continue

        if line.startswith("Scenario 2: High uncertainty"):
            scenario = "High uncertainty"
            results[scenario] = {}
            method = None
            continue

        if scenario is None:
            continue

        if line.startswith("---") and line.endswith("---"):
            method = line.strip("-").strip()
            in_total_results = False
            if method in METHODS:
                results[scenario].setdefault(method, {})
            continue

        if line == "TOTAL RESULTS:":
            in_total_results = True
            continue

        if line == "STATE SPECIFIC:":
            in_total_results = False
            continue

        if in_total_results and method in METHODS:
            match = re.match(
                r"^(including correct|only correct|correct parts)"
                r"\s+([\d.]+)%$",
                line,
            )
            if match:
                metric, value = match.groups()
                results[scenario][method][metric] = float(value) / 100

    return results


def create_excel(results):
    wb = Workbook()
    ws = wb.active
    ws.title = "Fusion comparison"

    # Grouped headers
    ws.merge_cells("A1:A2")
    ws["A1"] = "Fusion method"

    ws.merge_cells("B1:D1")
    ws["B1"] = "Low uncertainty"

    ws.merge_cells("E1:G1")
    ws["E1"] = "High uncertainty"

    for start_col in (2, 5):
        for offset, metric in enumerate(METRICS):
            ws.cell(row=2, column=start_col + offset, value=metric)

    # Results
    for row, method in enumerate(METHODS, start=3):
        ws.cell(row=row, column=1, value=method)

        for start_col, scenario in (
            (2, "Low uncertainty"),
            (5, "High uncertainty"),
        ):
            values = results.get(scenario, {}).get(method, {})

            for offset, metric in enumerate(METRICS):
                cell = ws.cell(
                    row=row,
                    column=start_col + offset,
                )

                value = values.get(metric)
                if value is not None:
                    cell.value = value
                    cell.number_format = "0.00%"

    # Formatting
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    subheader_fill = PatternFill("solid", fgColor="EAF2F8")
    border_side = Side(style="thin", color="B7C9D6")

    for row in ws.iter_rows(min_row=1, max_row=2, max_col=7):
        for cell in row:
            cell.font = Font(bold=True)
            cell.fill = (
                header_fill if cell.row == 1 else subheader_fill
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = Border(bottom=border_side)

    for row in ws.iter_rows(
        min_row=3,
        max_row=2 + len(METHODS),
        max_col=7,
    ):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="left" if cell.column == 1 else "center",
                vertical="center",
            )

    ws.column_dimensions["A"].width = 48
    for col in range(2, 8):
        ws.column_dimensions[get_column_letter(col)].width = 20

    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 34
    ws.freeze_panes = "B3"

    wb.save(OUTPUT_FILE)


def main():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Cannot find {INPUT_FILE}. "
            "Put the simulation log in this folder or change INPUT_FILE."
        )

    text = INPUT_FILE.read_text(encoding="utf-8")
    results = parse_output(text)
    create_excel(results)

    print(f"Excel file saved to: {OUTPUT_FILE.resolve()}")


if __name__ == "__main__":
    main()